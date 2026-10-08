from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from apps.assets.models import Asset, Installation, Procurement
from apps.core.models import Building, Centre, Division, Employee, Location, Place, Room, Section, SubSection, Supplier
from apps.core.services import audit
from apps.disposal.models import DisposalRecord
from apps.movements.models import AssetMovement
from .models import ImportBatch, ImportRow

ALIASES = {
    "ASSET_CODE": ["ASSET_CODE", "ASSET CODE", "RFID_CODE", "RFID CODE"],
    "INVENTORY_TYPE": ["INVENTORY_TYPE", "TYPE OF INVENTORY", "INVENTORY TYPE", "PIR/DIR", "PIR/DIR/IIR"],
    "MAIN_ITEM": ["MAIN_ITEM", "MAIN ITEM", "PARENT_ASSET", "PARENT ASSET", "MAIN ITEM TRANSACTION ID"],
    "SL_NO": ["SL_NO", "SL NO", "S.NO", "S NO"],
    "TRANSACTION_ID": ["TRANSACTION_ID", "TRANSACTION ID"],
    "NC_NO": ["NC_NO", "NC NO"],
    "NC_DATE": ["NC_DATE", "NC DATE"],
    "BRIEF_DESCRIPTION": ["BRIEF_DESCRIPTION", "BRIEF DESCRIPTION", "DESCRIPTION"],
    "SPECIFICATION": ["SPECIFICATION"],
    "MAKE": ["MAKE"],
    "MODEL": ["MODEL"],
    "ITEM_SL_NO": ["ITEM_SL_NO", "ITEM SL NO", "SERIAL NUMBER", "ITEM SERIAL NUMBER"],
    "QTY": ["QTY", "QUANTITY"],
    "AMOUNT": ["AMOUNT", "UNIT RATE", "AMOUNT/UNIT RATE", "AMOUNT / UNIT RATE"],
    "CURRENCY": ["CURRENCY", "BILL CURRENCY"],
    "PO_NO": ["PO_NO", "PO NO"],
    "PO_DATE": ["PO_DATE", "PO DATE"],
    "SUPPLIER_NAME": ["SUPPLIER_NAME", "SUPPLIER NAME"],
    "BILL_NO": ["BILL_NO", "BILL NO"],
    "BILL_DATE": ["BILL_DATE", "BILL DATE"],
    "BILL_VALUE": ["BILL_VALUE", "BILL VALUE"],
    "DRR_NO": ["DRR_NO", "DRR NO", "DRR NO."],
    "DRR_DATE": ["DRR_DATE", "DRR DATE"],
    "GRIN_NO": ["GRIN_NO", "GRIN NO", "GRIN NO."],
    "GRIN_DATE": ["GRIN_DATE", "GRIN DATE"],
    "STOCK_ENTRY_REFERENCE": ["STOCK_ENTRY_REFERENCE", "STOCK ENTRY REFERENCE"],
    "DATE_OF_INSTALLATION": ["DATE OF INSTALLATION", "DATE_OF_INSTALLATION"],
    "WARRANTY_PERIOD": ["WARRANTY PERIOD", "WARRANTY_PERIOD"],
    "STATUS_OF_ASSET": ["STATUS OF ASSET", "STATUS_OF_ASSET"],
    "LOG_BOOK_MAINTAINED": ["LOG_BOOK_MAINTAINED", "LOG BOOK MAINTAINED", "LOG BOOK"],
    "CENTRE": ["CENTRE", "CENTER"],
    "BUILDING": ["BUILDING", "BUILDING NAME", "BUILDING_NAME"],
    "ROOM_NO": ["ROOM_NO", "ROOM NO"],
    "ROOM_IN_CHARGE_NAME": ["ROOM_IN_CHARGE_NAME", "ROOM IN CHARGE NAME"],
    "EMP_ID": ["EMP_ID", "EMP ID"],
    "EMP_NAME": ["EMP_NAME", "EMP NAME"],
    "DIVISION": ["DIVISION"],
    "SECTION": ["SECTION"],
    "SUB_SECTION": ["SUB SECTION", "SUB_SECTION"],
    "MAIN_LOCATION": ["MAIN LOCATION", "MAIN_LOCATION"],
    "CURRENT_USER": ["CURRENT USER", "CURRENT_USER"],
    "TRANSFER": ["TRANSFER", "TRANSFER (YES/NO)"],
    "TRANSFER_FROM_NAME": ["TRANSFER_FROM_NAME", "TRANSFER FROM NAME"],
    "TRANSFER_FROM_ID": ["TRANSFER_FROM_ID", "TRANSFER FROM ID"],
    "TRANSFER_TO_ID": ["TRANSFER_TO_ID", "TRANSFER TO ID"],
    "TRANSFER_TO_NAME": ["TRANSFER_TO_NAME", "TRANSFER TO NAME"],
    "TRANS_VOUCHER_NO": ["TRANS_VOUCHER_NO", "TRANS VOUCHER NO", "TRANSFER VOUCHER NO"],
    "TRANSFER_VOUCHER_DATE": ["TRANSFER_VOUCHER_DATE", "TRANSFER VOUCHER DATE"],
    "RETURN": ["RETURN", "RETURN (YES/NO)"],
    "RETURNED_STOCK": ["RETURNED_STOCK", "RETURNED STOCK"],
    "RETURN_FROM_NAME": ["RETURN_FROM_NAME", "RETURN FROM NAME"],
    "RET_VOUCHER_NO": ["RET_VOUCHER_NO", "RET VOUCHER NO", "RETURN VOUCHER NO"],
    "RETURN_VOUCHER_DATE": ["RETURN_VOUCHER_DATE", "RETURN VOUCHER DATE"],
    "RETURN_CLAUSE": ["RETURN_CLAUSE", "RETURN CLAUSE"],
    "RETURN_FROM_DIVISION": ["RETURN_FROM_DIVISION", "RETURN FROM DIVISION"],
    "DISPOSAL_DETAILS_LOT_NAME": ["DISPOSAL_DETAILS_LOT NAME", "DISPOSAL DETAILS LOT NAME", "LOT NAME"],
    "DISPOSAL_FILE_NO": ["DISPOSAL FILE NO", "DISPOSAL_FILE_NO"],
    "AUCTION_ID": ["AUCTION ID", "AUCTION_ID"],
    "H1_BUYER": ["H1_BUYER", "H1 BUYER"],
    "EMD_AMOUNT": ["EMD AMOUNT", "EMD_AMOUNT"],
    "TOTAL_BOOK_VALUE": ["TOTAL BOOK VALUE", "TOTAL_BOOK_VALUE"],
    "FINANCIAL_YEAR": ["FINANCIAL YEAR", "FINANCIAL_YEAR"],
    "REALIZED_SALE_VALUE": ["REALIZED SALE VALUE", "REALIZED_SALE_VALUE"],
    "APPORTIONED_SALE_VALUE": ["APPORTIONED SALE VALUE", "APPORTIONED_SALE_VALUE"],
    "WRITE_OFF_OM_NO": ["WRITE OFF OM NO", "WRITE_OFF_OM_NO"],
    "WRITE_OFF_OM_DATE": ["WRITE OFF OM DATE", "WRITE_OFF_OM_DATE"],
    "PASSOUT_FOR": ["PASSOUT_FOR", "PASSOUT FOR"],
    "PASSOUT_TYPE": ["PASSOUT_TYPE", "PASSOUT TYPE"],
    "PASSOUT_NO": ["PASSOUT_NO", "PASSOUT NO"],
    "PASSOUT_DATE": ["PASSOUT_DATE", "PASSOUT DATE"],
    "REMARKS": ["REMARKS"],
}

DATE_FIELDS = {
    "NC_DATE", "PO_DATE", "BILL_DATE", "DATE_OF_INSTALLATION", "DRR_DATE", "GRIN_DATE",
    "TRANSFER_VOUCHER_DATE", "RETURN_VOUCHER_DATE", "WRITE_OFF_OM_DATE", "PASSOUT_DATE",
}
DECIMAL_FIELDS = {
    "AMOUNT", "BILL_VALUE", "EMD_AMOUNT", "TOTAL_BOOK_VALUE",
    "REALIZED_SALE_VALUE", "APPORTIONED_SALE_VALUE",
}


def clean_header(value):
    return " ".join(str(value or "").strip().upper().replace("_", " ").split())


def normalize_row(raw):
    normalized = {}
    source = {clean_header(k): v for k, v in raw.items()}
    for canonical, aliases in ALIASES.items():
        for alias in aliases:
            key = clean_header(alias)
            if key in source:
                normalized[canonical] = source[key]
                break
    return normalized


def iso_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def parse_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    try:
        # Excel date cells are stored in the preview as ISO datetimes (2026-07-01T00:00:00).
        return datetime.fromisoformat(str(value).strip()).date()
    except ValueError:
        pass
    raise ValueError(f"Invalid date: {value}")


def parse_decimal(value):
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).replace(",", "").strip())
    except InvalidOperation as exc:
        raise ValueError(f"Invalid amount: {value}") from exc


def parse_bool(value):
    return str(value or "").strip().lower() in {"yes", "y", "true", "1"}


def parse_status(value):
    text = str(value or "").strip().upper().replace(" ", "_")
    mapping = {
        "IN_STOCK": Asset.Status.IN_STOCK,
        "STOCK": Asset.Status.IN_STOCK,
        "ASSIGNED": Asset.Status.ASSIGNED,
        "ISSUED": Asset.Status.ASSIGNED,
        "RETURN_PENDING": Asset.Status.RETURN_PENDING,
        "WORKING": Asset.Status.ASSIGNED,
        "OUTSIDE": Asset.Status.OUTSIDE,
        "UNDER_DISPOSAL": Asset.Status.UNDER_DISPOSAL,
        "DISPOSED": Asset.Status.DISPOSED,
        "WRITTEN_OFF": Asset.Status.WRITTEN_OFF,
        "PASSED_OUT": Asset.Status.PASSED_OUT,
    }
    return mapping.get(text, Asset.Status.IN_STOCK)


def parse_optional_bool(value):
    text = str(value or "").strip().lower()
    if text in {"yes", "y", "true", "1"}:
        return True
    if text in {"no", "n", "false", "0"}:
        return False
    return None


def parse_working_status(value):
    text = " ".join(str(value or "").strip().lower().replace("_", " ").split())
    if text in {"working", "in working condition"}:
        return Installation.WorkingStatus.WORKING
    if text in {"not working", "non working", "non-working", "notworking"}:
        return Installation.WorkingStatus.NOT_WORKING
    return ""


def find_asset(reference):
    reference = str(reference or "").strip()
    if not reference:
        return None
    return Asset.objects.filter(transaction_id=reference).first() or Asset.objects.filter(asset_code=reference).first()


def validate_data(data, row_number, seen=None):
    """Validate one row. ``seen`` holds codes/transaction IDs of earlier rows so accessories can refer to them."""
    errors = []
    seen = set() if seen is None else seen
    if not data.get("BRIEF_DESCRIPTION"):
        errors.append("BRIEF_DESCRIPTION is required")
    for field in DATE_FIELDS:
        if data.get(field) not in (None, ""):
            try:
                parse_date(data[field])
            except ValueError as exc:
                errors.append(f"{field}: {exc}")
    for field in DECIMAL_FIELDS:
        if data.get(field) not in (None, ""):
            try:
                parse_decimal(data[field])
            except ValueError as exc:
                errors.append(f"{field}: {exc}")
    inventory_type = str(data.get("INVENTORY_TYPE") or "").strip().upper()
    if inventory_type and inventory_type not in Asset.InventoryType.values:
        errors.append("INVENTORY_TYPE must be PIR, DIR or IIR")
    currency = str(data.get("CURRENCY") or "").strip().upper()
    if currency and currency not in Procurement.Currency.values:
        errors.append("CURRENCY must be INR, USD, EUR or JPY")
    tx = str(data.get("TRANSACTION_ID") or "").strip()
    if tx and (tx in seen or Asset.objects.filter(transaction_id=tx).exists()):
        errors.append("TRANSACTION_ID already exists")
    code = str(data.get("ASSET_CODE") or "").strip()
    if code and (code in seen or Asset.objects.filter(asset_code=code).exists()):
        errors.append("RFID/ASSET_CODE already exists")
    main_item = str(data.get("MAIN_ITEM") or "").strip()
    if main_item and main_item not in seen and find_asset(main_item) is None:
        errors.append("MAIN_ITEM must be the transaction ID or RFID code of an existing asset or an earlier row")
    seen.update(value for value in (tx, code) if value)
    return errors


def validate_batch(batch):
    workbook = load_workbook(batch.uploaded_file.path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value or "").strip() for value in next(rows)]
    except StopIteration:
        headers = []
    batch.rows.all().delete()
    total = valid = invalid = 0
    seen = set()
    for excel_row, values in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in values):
            continue
        total += 1
        data = normalize_row(dict(zip(headers, values)))
        errors = validate_data(data, excel_row, seen)
        ImportRow.objects.create(
            batch=batch,
            row_number=excel_row,
            original_data={key: iso_value(value) for key, value in data.items()},
            validation_status=ImportRow.ValidationStatus.INVALID if errors else ImportRow.ValidationStatus.VALID,
            validation_errors=errors,
        )
        if errors:
            invalid += 1
        else:
            valid += 1
    batch.total_rows = total
    batch.valid_rows = valid
    batch.invalid_rows = invalid
    batch.status = ImportBatch.Status.VALIDATED
    batch.save(update_fields=["total_rows", "valid_rows", "invalid_rows", "status", "updated_at"])
    return batch


def get_division(name, prefix="DIV"):
    name = str(name or "").strip()
    if not name:
        return None
    existing = Division.objects.filter(name=name).first()
    if existing:
        return existing
    base = "".join(ch for ch in name.upper() if ch.isalnum())[:16] or prefix
    code = base
    counter = 1
    while Division.objects.filter(code=code).exists():
        counter += 1
        code = f"{base[:15]}{counter}"
    return Division.objects.create(name=name, code=code)


def get_organisation(data):
    division = section = subsection = None
    division_name = str(data.get("DIVISION") or "").strip()
    if division_name:
        division = get_division(division_name)
    section_name = str(data.get("SECTION") or "").strip()
    if section_name and division:
        section, _ = Section.objects.get_or_create(division=division, name=section_name)
    subsection_name = str(data.get("SUB_SECTION") or "").strip()
    if subsection_name and section:
        subsection, _ = SubSection.objects.get_or_create(section=section, name=subsection_name)
    return division, section, subsection


def get_employee(employee_id, name, division=None, section=None):
    employee_id = str(employee_id or "").strip()
    name = str(name or "").strip()
    if not employee_id and not name:
        return None
    if not employee_id:
        employee_id = "LEGACY-" + hashlib.sha1(f"{name}|{division_id(division)}".encode()).hexdigest()[:10].upper()
    employee, _ = Employee.objects.get_or_create(
        employee_id=employee_id,
        defaults={"name": name or employee_id, "division": division, "section": section},
    )
    return employee


def division_id(division):
    return getattr(division, "pk", None)


@transaction.atomic
def confirm_batch(batch, user):
    batch = ImportBatch.objects.select_for_update().get(pk=batch.pk)
    if batch.status != ImportBatch.Status.VALIDATED:
        raise ValueError("Batch must be validated first")
    if batch.invalid_rows:
        raise ValueError("Correct invalid rows before confirming the import")

    for row in batch.rows.select_for_update().filter(validation_status=ImportRow.ValidationStatus.VALID):
        data = row.original_data
        division, section, subsection = get_organisation(data)
        location = None
        location_name = str(data.get("MAIN_LOCATION") or "").strip()
        room_no = str(data.get("ROOM_NO") or "").strip()
        building = room = None
        building_name = str(data.get("BUILDING") or "").strip()
        if building_name:
            centre_name = str(data.get("CENTRE") or "").strip()
            centre = Centre.objects.get_or_create(name=centre_name)[0] if centre_name else None
            building = Building.objects.get_or_create(centre=centre, name=building_name)[0]
            if room_no:
                room = Room.objects.get_or_create(building=building, room_no=room_no)[0]
        if location_name or room_no:
            place = Place.objects.get_or_create(name=location_name or "Unspecified location")[0]
            location, _ = Location.objects.get_or_create(
                name=place.name,
                building=building,
                room_no=room_no,
                defaults={"place": place, "room": room, "division": division, "section": section, "sub_section": subsection},
            )
        custodian = get_employee(data.get("EMP_ID"), data.get("EMP_NAME"), division, section)
        supplier = None
        if data.get("SUPPLIER_NAME"):
            supplier, _ = Supplier.objects.get_or_create(name=str(data["SUPPLIER_NAME"]).strip())

        # SL No. and (when absent) the transaction ID are generated automatically.
        asset = Asset.objects.create(
            asset_code=str(data.get("ASSET_CODE") or "").strip() or None,
            inventory_type=str(data.get("INVENTORY_TYPE") or "").strip().upper(),
            parent_asset=find_asset(data.get("MAIN_ITEM")),
            division=division,
            transaction_id=str(data.get("TRANSACTION_ID") or "").strip(),
            nc_no=str(data.get("NC_NO") or "").strip(),
            nc_date=parse_date(data.get("NC_DATE")),
            brief_description=str(data["BRIEF_DESCRIPTION"]).strip(),
            specification=str(data.get("SPECIFICATION") or "").strip(),
            make=str(data.get("MAKE") or "").strip(),
            model=str(data.get("MODEL") or "").strip(),
            item_sl_no=str(data.get("ITEM_SL_NO") or "").strip(),
            current_status=parse_status(data.get("STATUS_OF_ASSET")),
            current_location=location,
            current_custodian=custodian,
            room_in_charge_name=str(data.get("ROOM_IN_CHARGE_NAME") or "").strip(),
            current_user=str(data.get("CURRENT_USER") or "").strip(),
            remarks=str(data.get("REMARKS") or "").strip(),
        )
        Procurement.objects.create(
            asset=asset,
            qty=int(float(data.get("QTY") or 1)),
            currency=str(data.get("CURRENCY") or "").strip().upper() or Procurement.Currency.INR,
            amount=parse_decimal(data.get("AMOUNT")),
            po_no=str(data.get("PO_NO") or ""),
            po_date=parse_date(data.get("PO_DATE")),
            supplier=supplier,
            bill_no=str(data.get("BILL_NO") or ""),
            bill_date=parse_date(data.get("BILL_DATE")),
            bill_value=parse_decimal(data.get("BILL_VALUE")),
            drr_no=str(data.get("DRR_NO") or ""),
            drr_date=parse_date(data.get("DRR_DATE")),
            grin_no=str(data.get("GRIN_NO") or ""),
            grin_date=parse_date(data.get("GRIN_DATE")),
        )
        warranty = data.get("WARRANTY_PERIOD")
        try:
            warranty = int(float(warranty)) if warranty not in (None, "") else None
        except (TypeError, ValueError):
            warranty = None
        Installation.objects.create(
            asset=asset,
            stock_entry_reference=str(data.get("STOCK_ENTRY_REFERENCE") or ""),
            date_of_installation=parse_date(data.get("DATE_OF_INSTALLATION")),
            warranty_period_months=warranty,
            status_of_asset=parse_working_status(data.get("STATUS_OF_ASSET")),
            log_book_maintained=parse_optional_bool(data.get("LOG_BOOK_MAINTAINED")),
        )

        if parse_bool(data.get("TRANSFER")):
            transfer_to = get_employee(data.get("TRANSFER_TO_ID"), data.get("TRANSFER_TO_NAME"), division, section)
            AssetMovement.objects.create(
                asset=asset,
                movement_type=AssetMovement.Type.TRANSFER,
                from_location=location,
                to_location=location,
                from_employee=get_employee(data.get("TRANSFER_FROM_ID"), data.get("TRANSFER_FROM_NAME"), division, section),
                to_employee=transfer_to,
                voucher_number=str(data.get("TRANS_VOUCHER_NO") or ""),
                movement_date=parse_date(data.get("TRANSFER_VOUCHER_DATE")) or timezone.localdate(),
                approved_by=user,
                transfer=True,
                transfer_from_name=str(data.get("TRANSFER_FROM_NAME") or ""),
                transfer_from_id=str(data.get("TRANSFER_FROM_ID") or ""),
                transfer_to_id=str(data.get("TRANSFER_TO_ID") or ""),
                transfer_to_name=str(data.get("TRANSFER_TO_NAME") or ""),
                transfer_voucher_date=parse_date(data.get("TRANSFER_VOUCHER_DATE")),
                remarks="Imported legacy transfer",
            )
            if transfer_to:
                asset.current_custodian = transfer_to
                asset.current_status = Asset.Status.ASSIGNED

        if parse_bool(data.get("RETURN")) or parse_bool(data.get("RETURNED_STOCK")):
            return_division_name = str(data.get("RETURN_FROM_DIVISION") or "").strip()
            return_division = None
            if return_division_name:
                return_division = get_division(return_division_name, prefix="RET")
            AssetMovement.objects.create(
                asset=asset,
                movement_type=AssetMovement.Type.RETURN,
                from_location=location,
                to_location=location,
                from_employee=asset.current_custodian,
                voucher_number=str(data.get("RET_VOUCHER_NO") or ""),
                movement_date=parse_date(data.get("RETURN_VOUCHER_DATE")) or timezone.localdate(),
                approved_by=user,
                returned_stock=True,
                return_from_name=str(data.get("RETURN_FROM_NAME") or ""),
                return_voucher_date=parse_date(data.get("RETURN_VOUCHER_DATE")),
                return_clause=str(data.get("RETURN_CLAUSE") or ""),
                return_from_division=return_division,
                remarks="Imported legacy return",
            )
            if location:
                location.is_stock_location = True
                location.save(update_fields=["is_stock_location", "updated_at"])
            asset.current_custodian = None
            asset.current_status = Asset.Status.IN_STOCK

        disposal_values = [
            data.get("DISPOSAL_DETAILS_LOT_NAME"), data.get("DISPOSAL_FILE_NO"), data.get("AUCTION_ID"),
            data.get("WRITE_OFF_OM_NO"), data.get("PASSOUT_NO"), data.get("H1_BUYER"),
        ]
        if any(value not in (None, "") for value in disposal_values):
            disposal_status = DisposalRecord.Status.PROPOSED
            if data.get("PASSOUT_NO") or data.get("PASSOUT_DATE"):
                disposal_status = DisposalRecord.Status.PASSED_OUT
                asset.current_status = Asset.Status.PASSED_OUT
            elif data.get("WRITE_OFF_OM_NO") or data.get("WRITE_OFF_OM_DATE"):
                disposal_status = DisposalRecord.Status.WRITTEN_OFF
                asset.current_status = Asset.Status.WRITTEN_OFF
            elif data.get("AUCTION_ID") or data.get("REALIZED_SALE_VALUE"):
                disposal_status = DisposalRecord.Status.AUCTIONED
                asset.current_status = Asset.Status.DISPOSED
            else:
                asset.current_status = Asset.Status.UNDER_DISPOSAL
            DisposalRecord.objects.create(
                asset=asset,
                status=disposal_status,
                lot_name=str(data.get("DISPOSAL_DETAILS_LOT_NAME") or ""),
                disposal_file_no=str(data.get("DISPOSAL_FILE_NO") or ""),
                auction_id=str(data.get("AUCTION_ID") or ""),
                h1_buyer=str(data.get("H1_BUYER") or ""),
                emd_amount=parse_decimal(data.get("EMD_AMOUNT")),
                total_book_value=parse_decimal(data.get("TOTAL_BOOK_VALUE")),
                financial_year=str(data.get("FINANCIAL_YEAR") or ""),
                realized_sale_value=parse_decimal(data.get("REALIZED_SALE_VALUE")),
                apportioned_sale_value=parse_decimal(data.get("APPORTIONED_SALE_VALUE")),
                write_off_om_no=str(data.get("WRITE_OFF_OM_NO") or ""),
                write_off_om_date=parse_date(data.get("WRITE_OFF_OM_DATE")),
                passout_for=str(data.get("PASSOUT_FOR") or ""),
                passout_type=str(data.get("PASSOUT_TYPE") or ""),
                passout_no=str(data.get("PASSOUT_NO") or ""),
                passout_date=parse_date(data.get("PASSOUT_DATE")),
                approved_by=user,
                remarks="Imported legacy disposal/write-off record",
            )

        asset.save(update_fields=["current_custodian", "current_status", "updated_at"])
        row.created_asset = asset
        row.validation_status = ImportRow.ValidationStatus.IMPORTED
        row.save(update_fields=["created_asset", "validation_status", "updated_at"])

    batch.status = ImportBatch.Status.COMPLETED
    batch.completed_at = timezone.now()
    batch.save(update_fields=["status", "completed_at", "updated_at"])
    audit(actor=user, action="IMPORT", obj=batch, description=f"Imported {batch.valid_rows} assets from Excel")
    return batch

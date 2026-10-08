import argparse
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from e2e.run_demo import DemoReport, demo_environment, is_local_url, make_import_workbook


class RunnerUnitTests(unittest.TestCase):
    def test_local_url_guard(self):
        self.assertTrue(is_local_url("http://127.0.0.1:8000"))
        self.assertTrue(is_local_url("http://localhost:8000"))
        self.assertFalse(is_local_url("https://example.com"))

    def test_demo_environment_uses_isolated_sqlite_and_removes_mysql(self):
        args = argparse.Namespace(use_configured_database=False)
        with tempfile.TemporaryDirectory() as temp_dir, patch.dict(
            os.environ,
            {"MYSQL_DATABASE": "production", "MYSQL_HOST": "db.example"},
            clear=False,
        ):
            env, database_path = demo_environment(args, Path(temp_dir))
        self.assertIsNotNone(database_path)
        self.assertEqual(env["SQLITE_DATABASE_PATH"], str(database_path))
        self.assertNotIn("MYSQL_DATABASE", env)
        self.assertNotIn("MYSQL_HOST", env)
        self.assertEqual(env["QR_PAGE_PUBLIC"], "False")

    def test_generated_import_workbook_contains_two_unique_demo_rows(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "demo.xlsx"
            make_import_workbook(path)
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0][0:3], ("ASSET_CODE", "TRANSACTION_ID", "BRIEF_DESCRIPTION"))
        self.assertEqual(rows[1][0], "PWDEMO-IMP-001")
        self.assertEqual(rows[2][0], "PWDEMO-IMP-002")
        self.assertNotEqual(rows[1][1], rows[2][1])

    def test_report_is_machine_readable(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "report.json"
            report = DemoReport(status="passed")
            report.save(path)
            html_path = Path(temp_dir) / "report.html"
            report.save_html(html_path)
            data = json.loads(path.read_text(encoding="utf-8"))
            html_text = html_path.read_text(encoding="utf-8")
        self.assertEqual(data["status"], "passed")
        self.assertIn("duration_seconds", data)
        self.assertIn("Asset-management guided feature tour", html_text)


    def test_required_ui_contract_is_present(self):
        root = Path(__file__).resolve().parents[1]
        contracts = {
            "templates/base.html": ["Assets", "Verification", "Movements", "Gate Passes", "Disposal", "Reports", "Administration"],
            "templates/assets/asset_detail.html": ["QR PNG", "Open QR page", "Movement history", "Procurement and installation"],
            "templates/movements/movement_list.html": ["Transfer", "Return to stock"],
            "templates/gatepasses/gatepass_detail.html": ["Division approve", "Stores approve", "Reject", "Mark outward", "Mark inward"],
            "templates/imports/batch_detail.html": ["Confirm import", "valid", "invalid"],
            "templates/reporting/dashboard.html": ["Acquisition value", "Overdue temporary gate passes", "Download asset register CSV"],
        }
        for relative_path, expected_text in contracts.items():
            content = (root / relative_path).read_text(encoding="utf-8")
            for text in expected_text:
                self.assertIn(text, content, f"{text!r} missing from {relative_path}")


if __name__ == "__main__":
    unittest.main()
